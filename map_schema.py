#!/usr/bin/env python3
"""
map_schema.py - Map ERP / MES exports with German or English headers into the DES schema template.

Usage:
    python map_schema.py --input <file or folder> --output out.xlsx
                         [--schema schema/Schema_Template_EMPTY.xlsx] [--synonyms synonyms.yaml]
                         [--config rules.yaml] [--answers answers.yaml] [--no-input]

Reads .csv/.txt, .xls/.xlsx/.xlsm (every sheet) and .accdb/.mdb (every table, Windows Access driver).
Headers are matched against synonyms.yaml. Certain matches are assigned, uncertain ones become questions,
missing columns are listed (required first) with two choices: give a file and column, or fill by hand.
Every answer is saved to an answers file, so a second run with --answers asks nothing already answered.
Runs locally: no network, no AI calls.
"""

import argparse
import difflib
import os
import re
import sys
import unicodedata
from collections import defaultdict

import openpyxl
import pandas as pd
import yaml

HERE = os.path.dirname(os.path.abspath(__file__))
DATA_EXT = (".csv", ".txt", ".xls", ".xlsx", ".xlsm", ".accdb", ".mdb")
SKIP_SHEETS = {"README", "Overview", "Dictionary", "Scenarios", "Config"}
TEXT_COLS = {"resource_group", "process_name", "required_skill", "scrap_or_rework", "setup_family", "setup_family_current",
             "worker_needed_during", "days", "start_time", "end_time", "break_start", "break_end", "resource_type",
             "valid_from", "valid_to", "order_type", "from_setup_family", "to_setup_family", "can_setup", "can_operate",
             "can_maintain"}
UNIT_WORDS = {"min": ["min", "mins", "minute", "minutes", "minuten"],
              "hours": ["h", "hr", "hrs", "hour", "hours", "std", "stunde", "stunden"],
              "sec": ["s", "sec", "secs", "second", "seconds", "sekunde", "sekunden"]}
TO_MIN = {"min": 1.0, "hours": 60.0, "sec": 1 / 60}

ANSWERS = {}
NO_INPUT = False
LOG = []


def say(msg=""):
    print(msg)
    LOG.append(msg)


# ── Text helpers ───────────────────────────────────────────────────────────────

def norm(h):
    s = str(h).replace("ß", "ss").lower()
    for a, b in (("ä", "ae"), ("ö", "oe"), ("ü", "ue")):
        s = s.replace(a, b)
    s = re.sub(r"\[.*?\]|\(.*?\)", " ", s)
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def is_text_col(col):
    return col.endswith("_id") or col.endswith("_name") or col in TEXT_COLS


def to_number(v):
    s = str(v).strip().replace(" ", "")
    if s == "":
        return None
    if re.fullmatch(r"-?\d{1,3}(\.\d{3})+,\d+", s):
        s = s.replace(".", "")
    s = s.replace(",", ".")
    try:
        f = float(s)
    except ValueError:
        return None
    return int(f) if f.is_integer() else f


# ── Questions (answered once, replayed from the answers file) ─────────────────

def ask(key, prompt, options):
    """options: list of (value, label). Returns the chosen value; None when unanswered in --no-input mode."""
    if key in ANSWERS:
        return ANSWERS[key]
    if NO_INPUT:
        return None
    print(f"\n? {prompt}")
    for i, (_, label) in enumerate(options, 1):
        print(f"   [{i}] {label}")
    while True:
        raw = read_line("   choice: ")
        if raw is None:
            return None
        if raw.isdigit() and 1 <= int(raw) <= len(options):
            ANSWERS[key] = options[int(raw) - 1][0]
            return ANSWERS[key]
        print("   please type one of the numbers above")


def ask_text(key, prompt, default=""):
    if key in ANSWERS:
        return ANSWERS[key]
    if NO_INPUT:
        return None
    raw = read_line(f"\n? {prompt}" + (f" [{default}]" if default else "") + ": ")
    if raw is None:
        return None
    ANSWERS[key] = raw or default
    return ANSWERS[key]


def read_line(prompt):
    """input() that switches to --no-input mode when the input ends (Ctrl+Z / Ctrl+D or a closed pipe)."""
    global NO_INPUT
    try:
        return input(prompt).strip()
    except EOFError:
        NO_INPUT = True
        say("\n  (input ended: remaining questions stay open and are listed in the report)")
        return None


# ── Readers ────────────────────────────────────────────────────────────────────

def clean(df):
    df = df.fillna("").astype(str)
    return df.apply(lambda c: c.str.strip())


def header_row(raw):
    counts = [int((raw.iloc[i] != "").sum()) for i in range(min(10, len(raw)))]
    if not counts or max(counts) < 2:
        return None
    return next(i for i, n in enumerate(counts) if n >= 0.8 * max(counts))


def from_grid(raw):
    raw = clean(raw)
    h = header_row(raw)
    if h is None:
        return None
    cols, seen = [], defaultdict(int)
    for c in raw.iloc[h]:
        c = c or f"column_{len(cols) + 1}"
        seen[c] += 1
        cols.append(c if seen[c] == 1 else f"{c}_{seen[c]}")
    df = raw.iloc[h + 1:].copy()
    df.columns = cols
    df = df[(df != "").any(axis=1)]
    return df.reset_index(drop=True) if len(df) else None


def read_csv(path):
    for enc in ("utf-8-sig", "cp1252"):
        try:
            with open(path, encoding=enc) as fh:
                first = fh.readline()
            sep = max([";", ",", "\t", "|"], key=first.count)
            raw = pd.read_csv(path, sep=sep, dtype=str, encoding=enc, header=None, keep_default_na=False)
            return from_grid(raw)
        except UnicodeDecodeError:
            continue
    return None


def read_access(path):
    try:
        import pyodbc
    except ImportError:
        say(f"  ! {os.path.basename(path)}: reading Access files needs the pyodbc package (pip install pyodbc)")
        return {}
    drivers = [d for d in pyodbc.drivers() if "Access" in d and "accdb" in d]
    if not drivers:
        say(f"  ! {os.path.basename(path)}: no Microsoft Access ODBC driver on this computer; export the tables to .xlsx or .csv")
        return {}
    con = pyodbc.connect(f"DRIVER={{{drivers[0]}}};DBQ={os.path.abspath(path)};", autocommit=True)
    out = {}
    try:
        names = [t.table_name for t in con.cursor().tables(tableType="TABLE")]
        for name in names:
            try:
                df = access_table(con, name, None)
            except pyodbc.Error:
                cols = [c.column_name for c in con.cursor().columns(table=name)]
                good = [c for c in cols if access_column_ok(con, name, c)]
                say(f"  ! {os.path.basename(path)}:{name}: unreadable columns skipped: {sorted(set(cols) - set(good))}")
                df = access_table(con, name, good) if good else None
            if df is not None and len(df):
                out[name] = clean(df)
    finally:
        con.close()
    return out


def access_table(con, name, cols):
    sel = ", ".join(f"[{c}]" for c in cols) if cols else "*"
    cur = con.cursor().execute(f"SELECT {sel} FROM [{name}]")
    return pd.DataFrame.from_records([tuple(r) for r in cur.fetchall()], columns=[d[0] for d in cur.description])


def access_column_ok(con, name, col):
    try:
        con.cursor().execute(f"SELECT [{col}] FROM [{name}]").fetchall()
        return True
    except Exception:
        return False


def load_tables(inp):
    files = []
    if os.path.isdir(inp):
        for root, _, names in os.walk(inp):
            files += [os.path.join(root, n) for n in sorted(names) if n.lower().endswith(DATA_EXT)]
        base = inp
    else:
        files, base = [inp], os.path.dirname(inp)
    tables = {}
    for f in files:
        rel = os.path.relpath(f, base).replace("\\", "/")
        ext = os.path.splitext(f)[1].lower()
        try:
            if ext in (".csv", ".txt"):
                df = read_csv(f)
                if df is not None:
                    tables[rel] = df
            elif ext in (".accdb", ".mdb"):
                for name, df in read_access(f).items():
                    tables[f"{rel}:{name}"] = df
            else:
                engine = "xlrd" if ext == ".xls" else "openpyxl"
                for name, raw in pd.read_excel(f, sheet_name=None, header=None, dtype=str, engine=engine).items():
                    if name in SKIP_SHEETS:
                        continue
                    df = from_grid(raw)
                    if df is not None:
                        tables[f"{rel}:{name}"] = df
        except Exception as exc:
            say(f"  ! could not read {rel}: {exc}")
    return tables, base


# ── Template ───────────────────────────────────────────────────────────────────

def load_template(path):
    wb = openpyxl.load_workbook(path)
    sheets = {}
    for ws in wb.worksheets:
        if ws.title in SKIP_SHEETS:
            continue
        cols = [c.value for c in ws[2] if c.value]
        banner = str(ws["A1"].value or "")
        sheets[ws.title] = {"cols": cols, "required": "[REQUIRED]" in banner}
    required_cols, blank_rule = set(), {}
    if "Dictionary" in wb.sheetnames:
        ws = wb["Dictionary"]
        hdr = [c.value for c in ws[2]]
        i_s, i_c, i_r, i_b = (hdr.index(h) for h in ("sheet", "column", "required_optional", "if_blank_or_missing"))
        for row in ws.iter_rows(min_row=3, values_only=True):
            rule = str(row[i_b] or "")
            blank_rule[(row[i_s], row[i_c])] = rule
            if row[i_r] == "required" and ("skipped" in rule or "REJECTED" in rule):
                required_cols.add((row[i_s], row[i_c]))
    for s, info in sheets.items():
        info["blank"] = {c: blank_rule.get((s, c), "") for c in info["cols"]}
        sheet_level = len(info["cols"]) > 1 and len(set(info["blank"].values())) == 1
        info["req_cols"] = [c for c in info["cols"] if (s, c) in required_cols and (not sheet_level or c == info["cols"][0])]
    count = defaultdict(int)
    for info in sheets.values():
        for c in info["cols"]:
            count[c] += 1
    for info in sheets.values():
        info["distinct"] = [c for c in info["cols"] if count[c] == 1]
    return sheets


def load_synonyms(path):
    raw = yaml.safe_load(open(path, encoding="utf-8")) or {}
    syn = {}
    for col, d in raw.items():
        syn[col] = {"exact": {norm(x) for x in d.get("exact") or []} | {norm(col)},
                    "ambiguous": {norm(x) for x in d.get("ambiguous") or []}}
    return syn


def score(header, col, syn):
    """2 = certain, 1 = possible (becomes a question), 0 = no match."""
    s = syn.get(col, {"exact": {norm(col)}, "ambiguous": set()})
    n = norm(header)
    if not n:
        return 0
    if n in s["exact"]:
        return 2
    if n in s["ambiguous"]:
        return 1
    toks = set(n.split())
    for e in s["exact"]:
        et = set(e.split())
        if len(et) >= 2 and et <= toks:
            return 1
    if max(difflib.SequenceMatcher(None, n, e).ratio() for e in s["exact"]) >= 0.88:
        return 1
    return 0


# ── Config (optional rules, also the legacy extract.py format) ────────────────

LEGACY = {"part_id": "product_id", "part_name": "product_name", "routing_step": "step_number",
          "operation_name": "process_name", "cost_center": "resource_group", "calc_rate": "cycle_time_min"}


def load_rules(path):
    if not path:
        return {"columns": {}, "filters": {}, "skip_values": {}, "unit_rules": {}, "transforms": {}, "dedupe": None}
    cfg = yaml.safe_load(open(path, encoding="utf-8")) or {}
    rules = {"columns": dict(cfg.get("columns_map") or {}), "filters": dict(cfg.get("filters") or {}),
             "skip_values": dict(cfg.get("skip_values") or {}), "unit_rules": dict(cfg.get("unit_rules") or {}),
             "transforms": dict(cfg.get("transforms") or {}), "dedupe": cfg.get("dedupe")}
    legacy = cfg.get("columns") or {}
    for old, new in LEGACY.items():
        if old in legacy:
            rules["columns"].setdefault(new, legacy[old])
    if "cost_center" in legacy:
        rules["transforms"].setdefault("resource_group", {"prefix": "RG_"})
    if (cfg.get("site_filter") or {}).get("value") and "site" in legacy:
        rules["filters"].setdefault(legacy["site"], cfg["site_filter"]["value"])
    if cfg.get("skip_operation_names") and "operation_name" in legacy:
        rules["skip_values"].setdefault(legacy["operation_name"], cfg["skip_operation_names"])
    if cfg.get("zbs_rules") and "zbs" in legacy:
        z = cfg["zbs_rules"]
        rules["unit_rules"].setdefault("cycle_time_min", {
            "code_column": legacy["zbs"], "skip_codes": z.get("skip_values") or [],
            "codes": {str(k): v for k, v in (z.get("conversions") or {}).items()}})
    if "routing_plan_version" in legacy and not rules["dedupe"]:
        keys = [legacy[k] for k in ("part_id", "routing_plan", "routing_step") if k in legacy]
        rules["dedupe"] = {"keys": keys, "version": legacy["routing_plan_version"]}
    return rules


def apply_filters(name, df, rules):
    before = len(df)
    for col, val in rules["filters"].items():
        if col in df.columns:
            df = df[df[col] == str(val)]
    for col, vals in rules["skip_values"].items():
        if col in df.columns:
            df = df[~df[col].isin([str(v) for v in vals])]
    ur = rules["unit_rules"].get("cycle_time_min")
    if ur and ur["code_column"] in df.columns:
        skip = {str(v).strip().upper() for v in ur["skip_codes"]}
        df = df[~df[ur["code_column"]].str.upper().isin(skip)]
    d = rules["dedupe"]
    if d and all(k in df.columns for k in d["keys"] + [d["version"]]):
        df = df.assign(_v=pd.to_numeric(df[d["version"]].str.replace(",", "."), errors="coerce"))
        df = df.sort_values(d["keys"] + ["_v"], na_position="first").drop_duplicates(d["keys"], keep="last").drop(columns="_v")
    if len(df) != before:
        say(f"  rules on {name}: {before:,} -> {len(df):,} rows (filters, skipped values, latest version)")
    return df.reset_index(drop=True)


# ── Mapping ────────────────────────────────────────────────────────────────────

def best_cols(df, sheet, info, syn, rules):
    """Per template column: (certain header or None, [possible headers])."""
    out = {}
    for col in info["cols"]:
        forced = rules["columns"].get(f"{sheet}.{col}") or rules["columns"].get(col)
        if forced and forced in df.columns:
            out[col] = (forced, [])
            continue
        sc = {h: score(h, col, syn) for h in df.columns}
        sure = [h for h, v in sc.items() if v == 2]
        maybe = [h for h, v in sc.items() if v == 1]
        out[col] = (sure[0], []) if len(sure) == 1 else (None, sure + maybe)
    return out


def table_score(m, info):
    """(points, certain sheet-specific columns): certain match = 1 point, possible = 0.5."""
    pts = sum(1.0 if sure else 0.5 if maybe else 0 for sure, maybe in m.values())
    key = info["cols"][0]
    if not (m[key][0] or m[key][1]):
        return 0, 0
    distinct = [c for c in info["distinct"] if m[c][0] or m[c][1]]
    if not distinct or sum(1 for v in m.values() if v[0] or v[1]) < 2:
        return 0, 0
    return pts, sum(1 for c in info["distinct"] if m[c][0])


def choose_table(sheet, info, tables, syn, rules):
    cands = []
    for name, df in tables.items():
        m = best_cols(df, sheet, info, syn, rules)
        pts, sure = table_score(m, info)
        if pts:
            cands.append((pts, sure, name))
    cands.sort(key=lambda c: (c[0], c[1]), reverse=True)
    if not cands:
        return None
    top = cands[0]
    if top[1] and top[0] >= 2 and (len(cands) == 1 or top[0] >= cands[1][0] + 1):
        say(f"  {sheet:<22} <- {top[2]}  (certain, score {top[0]:g})")
        return top[2]
    opts = [(c[2], f"{c[2]}  (score {c[0]:g})") for c in cands[:4]] + [("-", "none of these: leave the sheet empty")]
    pick = ask(f"{sheet}|table", f"Which table holds the data for sheet '{sheet}'?", opts)
    if pick in (None, "-"):
        say(f"  {sheet:<22} <- (open question)" if pick is None else f"  {sheet:<22} <- none (your answer)")
        return None
    say(f"  {sheet:<22} <- {pick}  (your answer)")
    return pick


def assign_columns(sheet, info, df, syn, rules):
    m = best_cols(df, sheet, info, syn, rules)
    used = {sure for sure, _ in m.values() if sure}
    mapping = {}
    for col in info["cols"]:
        sure, maybe = m[col]
        if sure:
            mapping[col] = sure
            say(f"    {col:<32} = {sure}")
            continue
        maybe = [h for h in maybe if h not in used]
        if not maybe:
            continue
        opts = [(h, f"'{h}'  e.g. {first_values(df, h)}") for h in maybe[:6]] + [("-", "none of these")]
        pick = ask(f"{sheet}|{col}", f"Sheet {sheet}, column '{col}': is it one of these headers?", opts)
        if pick and pick != "-":
            mapping[col] = pick
            used.add(pick)
            say(f"    {col:<32} = {pick}  (your answer)")
        elif pick is None:
            say(f"    {col:<32} ? possible: {', '.join(maybe[:6])}  (open question)")
    return mapping


def first_values(df, h, n=3):
    vals = [v for v in df[h].tolist() if v][:n]
    return ", ".join(vals) if vals else "(empty)"


def unit_of(header):
    toks = norm(re.sub(r"[\[\]()]", " ", str(header))).split()
    for unit, words in UNIT_WORDS.items():
        if any(t in words for t in toks):
            return unit
    return None


def parse_dates(series):
    vals = series[series != ""]
    if vals.empty or vals.str.fullmatch(r"-?[\d.,]+").all():
        return None
    parsed = pd.to_datetime(vals, errors="coerce", dayfirst=vals.str.contains(r"^\d{1,2}\.\d{1,2}\.").any())
    return pd.to_datetime(series, errors="coerce", dayfirst=vals.str.contains(r"^\d{1,2}\.\d{1,2}\.").any()) \
        if parsed.notna().mean() >= 0.9 else None


def convert(sheet, col, header, df, rules):
    """Returns the column as a list of values in the template unit, or None when the unit stays open."""
    raw = df[header]
    target = "min" if col.endswith("_min") else "hours" if col.endswith("_hours") else None
    ur = rules["unit_rules"].get(col)
    if ur and ur["code_column"] in df.columns:
        out = []
        for v, code in zip(raw, df[ur["code_column"]]):
            num, rule = to_number(v), ur["codes"].get(str(code).strip())
            out.append(None if num is None else round(num / float(rule["divisor"]) * float(rule["multiplier"]), 6) if rule else num)
        unknown = sorted({str(c) for c in df[ur["code_column"]] if str(c).strip() not in ur["codes"]})
        say(f"    {col}: converted by rate code column '{ur['code_column']}'" + (f"; unknown codes kept raw: {unknown}" if unknown else ""))
        return out
    if target is None:
        return list(raw)
    dates = parse_dates(raw)
    if dates is not None:
        if target != "min":
            say(f"    {col}: '{header}' holds dates; left for you to fill")
            return None
        start = ask_text("time0", f"'{header}' holds dates. Simulation start (time 0 for every date column) as YYYY-MM-DD HH:MM",
                         dates.min().normalize().strftime("%Y-%m-%d %H:%M"))
        if not start:
            return None
        t0 = pd.Timestamp(start)
        say(f"    {col}: dates -> minutes after {t0:%Y-%m-%d %H:%M}")
        return [None if pd.isna(d) else round((d - t0).total_seconds() / 60, 3) for d in dates]
    unit = unit_of(header)
    if unit:
        factor = TO_MIN[unit] / TO_MIN[target]
    else:
        opts = [(1.0, f"already in {target}")]
        opts += [(TO_MIN[u] / TO_MIN[target], f"in {u}") for u in TO_MIN if u != target]
        opts += [("other", "other unit: I give a factor")]
        factor = ask(f"{sheet}|{col}|unit", f"Unit of '{header}' (for {col}, e.g. {first_values(df, header)})?", opts)
        if factor == "other":
            f = ask_text(f"{sheet}|{col}|factor", f"Multiply '{header}' by which factor to get {target}?")
            factor = to_number(f) if f else None
            ANSWERS[f"{sheet}|{col}|unit"] = factor
        if factor is None:
            say(f"    {col}: unit of '{header}' open; left for you to fill")
            return None
    if factor != 1.0:
        say(f"    {col}: '{header}' x {factor:g} -> {target}")
    nums = [to_number(v) for v in raw]
    return [None if n is None else round(n * factor, 6) for n in nums]


def apply_transforms(sheet, col, values, rules):
    t = rules["transforms"].get(f"{sheet}.{col}") or rules["transforms"].get(col)
    if not t:
        return values
    out = []
    for v in values:
        s = "" if v is None else str(v)
        if s and t.get("strip_prefix") and s.startswith(t["strip_prefix"]):
            s = s[len(t["strip_prefix"]):]
        if s and t.get("prefix"):
            s = t["prefix"] + s
        out.append(s)
    return out


def fill_missing(sheet, info, df, mapping, tables, base):
    """Asks for each unfilled column: give a file and column, or fill by hand. Returns {col: values}."""
    extra = {}
    todo = [c for c in info["req_cols"] if c not in mapping] + [c for c in info["cols"] if c not in mapping and c not in info["req_cols"]]
    if not todo:
        return extra
    labels = [f"{c} ({'REQUIRED' if c in info['req_cols'] else 'optional'})" for c in todo]
    pick = ask(f"{sheet}|missing", f"Sheet {sheet}: no column found for {', '.join(labels)}. What now?",
               [("hand", "fill them by hand later"), ("source", "I give the file and column for some of them")])
    if pick != "source":
        return extra
    for col in todo:
        kind = "REQUIRED" if col in info["req_cols"] else "optional"
        sub = ask(f"{sheet}|{col}|missing", f"'{col}' ({kind}):",
                  [("hand", "by hand later"), ("source", "I give the file and column")])
        if sub != "source":
            continue
        vals = column_from_source(sheet, col, df, tables, base)
        if vals is not None:
            extra[col] = vals
    return extra


def column_from_source(sheet, col, df, tables, base):
    src = ask_text(f"{sheet}|{col}|file", f"File or table for '{col}' (path, or a table name from the list above)")
    if not src:
        return None
    other = tables.get(src)
    if other is None:
        path = src if os.path.isabs(src) else os.path.join(base, src)
        loaded, _ = load_tables(path) if os.path.exists(path) else ({}, None)
        other = next(iter(loaded.values()), None)
        if other is None:
            say(f"    {col}: could not read '{src}'")
            return None
    hdr = ask_text(f"{sheet}|{col}|column", f"Column in '{src}' for '{col}' (headers: {', '.join(other.columns[:15])})")
    if not hdr or hdr not in other.columns:
        say(f"    {col}: column '{hdr}' not in '{src}'")
        return None
    if other is df or other.equals(df):
        return list(df[hdr])
    shared = [h for h in df.columns if h in other.columns]
    on = ask(f"{sheet}|{col}|join", f"Link the rows of both tables by which shared column?",
             [(h, h) for h in shared] + [("-", "no shared column: cancel")]) if shared else None
    if not on or on == "-":
        say(f"    {col}: no column links '{src}' to the {sheet} table; left for you to fill")
        return None
    lookup, repeated = {}, set()
    for k, v in zip(other[on], other[hdr]):
        if k in lookup:
            repeated.add(k)
        else:
            lookup[k] = v
    say(f"    {col} = {src}:{hdr}  (linked by '{on}')")
    if repeated:
        say(f"    ! {len(repeated)} values of '{on}' appear several times in {src}; the first row was used (no sums or averages)")
    return [lookup.get(k, "") for k in df[on]]


# ── Build, write, report ───────────────────────────────────────────────────────

def build_sheet(sheet, info, df, mapping, extra, rules):
    data, open_cols = {}, []
    for col in info["cols"]:
        if col in mapping:
            vals = convert(sheet, col, mapping[col], df, rules)
        elif col in extra:
            vals = extra[col]
        else:
            continue
        if vals is None:
            open_cols.append(col)
            continue
        data[col] = apply_transforms(sheet, col, vals, rules)
    if not data:
        return [], open_cols
    rows = [dict(zip(data, r)) for r in zip(*data.values())]
    key = info["cols"][0]
    rows = [r for r in rows if r.get(key) not in (None, "")]
    unique_key = sheet in ("Products", "Machines", "Workers", "Shifts")
    if unique_key or len(data) < len(info["cols"]):
        seen, kept = set(), []
        for r in rows:
            sig = r.get(key) if unique_key else tuple(r.items())
            if sig not in seen:
                seen.add(sig)
                kept.append(r)
        if len(kept) != len(rows):
            say(f"    {len(rows) - len(kept):,} duplicate rows dropped ({'same ' + key if unique_key else 'identical'})")
        rows = kept
    return rows, open_cols


def write_workbook(template, output, results):
    wb = openpyxl.load_workbook(template)
    for sheet, rows in results.items():
        ws = wb[sheet]
        hdr = [c.value for c in ws[2]]
        for r in ws.iter_rows(min_row=3):
            for c in r:
                c.value = None
        for i, row in enumerate(rows, 3):
            for j, col in enumerate(hdr, 1):
                v = row.get(col)
                if v in (None, ""):
                    continue
                num = None if is_text_col(col) else to_number(v)
                ws.cell(row=i, column=j, value=num if num is not None else v)
    if "Config" in wb.sheetnames:
        ws = wb["Config"]
        for r in ws.iter_rows(min_row=3):
            if r[0].value == "scenario_id":
                r[1].value = os.path.splitext(os.path.basename(output))[0]
    wb.save(output)


def main():
    global NO_INPUT
    ap = argparse.ArgumentParser(description="Map ERP/MES exports into the DES schema template.")
    ap.add_argument("--input", required=True, help="export file (.csv .xls .xlsx .accdb) or a folder of them")
    ap.add_argument("--output", required=True, help="workbook to write (.xlsx)")
    ap.add_argument("--schema", default=os.path.join(HERE, "schema", "Schema_Template_EMPTY.xlsx"))
    ap.add_argument("--synonyms", default=os.path.join(HERE, "synonyms.yaml"))
    ap.add_argument("--config", help="optional rules: filters, rate-code units, prefixes, latest version")
    ap.add_argument("--answers", help="answers file of an earlier run (default: <output>_answers.yaml if it exists)")
    ap.add_argument("--no-input", action="store_true", help="ask nothing; open questions are listed in the report")
    a = ap.parse_args()
    NO_INPUT = a.no_input
    stem = os.path.splitext(a.output)[0]
    ans_path = a.answers or f"{stem}_answers.yaml"
    if os.path.exists(ans_path):
        ANSWERS.update(yaml.safe_load(open(ans_path, encoding="utf-8")) or {})
        say(f"Answers loaded from {ans_path} ({len(ANSWERS)})")

    sheets = load_template(a.schema)
    syn = load_synonyms(a.synonyms)
    rules = load_rules(a.config)

    say(f"\n[1/4] Reading {a.input}")
    tables, base = load_tables(a.input)
    if not tables:
        sys.exit("No readable table found.")
    tables = {n: apply_filters(n, df, rules) for n, df in tables.items()}
    for n, df in tables.items():
        say(f"  {n:<45} {len(df):>8,} rows  {len(df.columns):>3} columns")

    say("\n[2/4] Tables -> sheets")
    chosen = {s: choose_table(s, info, tables, syn, rules) for s, info in sheets.items()}

    say("\n[3/4] Columns")
    results, report = {}, {}
    for s, info in sheets.items():
        t = chosen[s]
        if not t:
            continue
        df = tables[t]
        say(f"  {s} <- {t}")
        mapping = assign_columns(s, info, df, syn, rules)
        extra = fill_missing(s, info, df, mapping, tables, base)
        rows, open_cols = build_sheet(s, info, df, mapping, extra, rules)
        results[s] = rows
        report[s] = {"table": t, "rows": len(rows), "filled": [c for c in info["cols"] if c in mapping or c in extra],
                     "open": open_cols}

    for s, info in sheets.items():
        if s not in results and info["required"]:
            pick = ask(f"{s}|missing_sheet", f"No table found for the required sheet '{s}'. What now?",
                       [("hand", "fill it by hand later"), ("source", "I give the file")])
            if pick == "source":
                src = ask_text(f"{s}|sheet_file", f"File or table for sheet '{s}'")
                df = tables.get(src)
                if df is None and src and os.path.exists(os.path.join(base, src)):
                    df = next(iter(load_tables(os.path.join(base, src))[0].values()), None)
                if df is not None:
                    mapping = assign_columns(s, info, df, syn, rules)
                    rows, open_cols = build_sheet(s, info, df, mapping, {}, rules)
                    results[s] = rows
                    report[s] = {"table": src, "rows": len(rows), "filled": list(mapping), "open": open_cols}

    write_workbook(a.schema, a.output, results)
    with open(ans_path, "w", encoding="utf-8") as fh:
        yaml.safe_dump(ANSWERS, fh, allow_unicode=True, sort_keys=True)

    say(f"\n[4/4] Written {a.output}   (answers saved to {ans_path})")
    say("\nREQUIRED, still empty:")
    for s, info in sheets.items():
        if info["required"] and s not in results:
            say(f"  {s}: whole sheet")
        elif s in report:
            miss = [c for c in info["req_cols"] if c not in report[s]["filled"] or c in report[s]["open"]]
            if miss:
                say(f"  {s}: {', '.join(miss)}")
    say("\nOPTIONAL, still empty (the model uses the default from the Dictionary sheet, or the module stays off):")
    for s, info in sheets.items():
        if not info["required"] and s not in results:
            say(f"  {s}: whole sheet (module off)")
        elif s in report:
            miss = [c for c in info["cols"] if c not in info["req_cols"] and (c not in report[s]["filled"] or c in report[s]["open"])]
            for c in miss:
                rule = info["blank"].get(c, "")
                say(f"  {s}.{c}" + (f"  -> {rule}" if rule.startswith("DEFAULT") else ""))
    say("\nAlways check by hand: Config (run length, default_machine_shift, modes); units and meaning of every column (Dictionary).")
    with open(f"{stem}_report.txt", "w", encoding="utf-8") as fh:
        fh.write("\n".join(LOG) + "\n")


if __name__ == "__main__":
    main()
